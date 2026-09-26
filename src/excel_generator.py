"""
excel_generator.py
------------------
Generates a formatted Excel cost-report workbook from structured row data.

Single responsibility: accept the list[dict] produced by data_provider.get_report_data()
and return a ready-to-upload BytesIO buffer containing the workbook.

This module has **no AWS dependencies** and no knowledge of where the workbook
will be stored.  It can therefore be unit-tested in isolation.

Column layout (11 columns, A–K):
    A: Item
    B: Location
    C: Jan 2026
    D: Feb 2026
    E: Mar 2026
    F: Apr 2026
    G: May 2026
    H: Jun 2026
    I: Jul 2026
    J: Aug 2026
    K: Sep 2026

Dict-key → column-header mapping is defined in the module-level constants
COLUMN_HEADERS and COLUMN_KEYS so both lists stay in sync automatically.
"""

from io import BytesIO

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

# ---------------------------------------------------------------------------
# Module-level constants — single source of truth for column order
# ---------------------------------------------------------------------------

# Display headers written to row 1 of the worksheet (order matters).
COLUMN_HEADERS: list[str] = [
    "Item",
    "Location",
    "Jan 2026",
    "Feb 2026",
    "Mar 2026",
    "Apr 2026",
    "May 2026",
    "Jun 2026",
    "Jul 2026",
    "Aug 2026",
    "Sep 2026",
]

# Dict keys from the data-provider schema that map to the columns above.
# Index i in COLUMN_KEYS corresponds to index i in COLUMN_HEADERS.
COLUMN_KEYS: list[str] = [
    "item",
    "location",
    "jan_2026",
    "feb_2026",
    "mar_2026",
    "apr_2026",
    "may_2026",
    "jun_2026",
    "jul_2026",
    "aug_2026",
    "sep_2026",
]

# Number format applied to every monthly cost cell (columns C–K).
COST_NUMBER_FORMAT: str = "#,##0.00"

# Excel table style — medium weight with row banding.
TABLE_STYLE: str = "TableStyleMedium9"

# Column widths (in character units).
COLUMN_WIDTHS: dict[str, int] = {
    "A": 28,  # Item
    "B": 14,  # Location
    "C": 12,  # Jan 2026
    "D": 12,  # Feb 2026
    "E": 12,  # Mar 2026
    "F": 12,  # Apr 2026
    "G": 12,  # May 2026
    "H": 12,  # Jun 2026
    "I": 12,  # Jul 2026
    "J": 12,  # Aug 2026
    "K": 12,  # Sep 2026
}


def generate_excel(rows: list[dict]) -> BytesIO:
    """Generate a formatted Excel cost-report workbook from structured row data.

    Assumes each dict in ``rows`` conforms to the stable schema defined in
    data_provider.py and design.md §4.1.  The function does not validate the
    input schema at runtime; callers are responsible for providing correctly
    shaped data.

    Parameters
    ----------
    rows : list[dict]
        Cost-report rows, each containing the keys defined in ``COLUMN_KEYS``.
        Typically the output of ``data_provider.get_report_data()``.

    Returns
    -------
    BytesIO
        A seeked-to-zero in-memory buffer containing a valid ``.xlsx`` workbook.
        Pass ``buffer.getvalue()`` or the buffer itself directly to boto3 for
        S3 upload — no disk I/O required.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cost Report"

    # ------------------------------------------------------------------
    # 1. Write header row (row 1) with bold font
    # ------------------------------------------------------------------
    bold_font = Font(bold=True)
    for col_idx, header in enumerate(COLUMN_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = bold_font

    # ------------------------------------------------------------------
    # 2. Write data rows (rows 2 … n+1)
    # ------------------------------------------------------------------
    for row_idx, row_data in enumerate(rows, start=2):
        for col_idx, key in enumerate(COLUMN_KEYS, start=1):
            ws.cell(row=row_idx, column=col_idx, value=row_data[key])

    # ------------------------------------------------------------------
    # 3. Apply numeric format to cost columns (C=3 … K=11) for all data rows
    # ------------------------------------------------------------------
    num_data_rows = len(rows)
    for row_idx in range(2, num_data_rows + 2):          # rows 2 … n+1
        for col_idx in range(3, len(COLUMN_HEADERS) + 1): # cols 3 … 11
            ws.cell(row=row_idx, column=col_idx).number_format = COST_NUMBER_FORMAT

    # ------------------------------------------------------------------
    # 4. Excel Table over the full data range (A1:K{last_row})
    # ------------------------------------------------------------------
    last_row = num_data_rows + 1  # +1 for the header row
    last_col_letter = openpyxl.utils.get_column_letter(len(COLUMN_HEADERS))
    table_ref = f"A1:{last_col_letter}{last_row}"

    table = Table(displayName="CostReport", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name=TABLE_STYLE,
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    # ------------------------------------------------------------------
    # 5. Freeze header row (everything above row 2 stays visible on scroll)
    # ------------------------------------------------------------------
    ws.freeze_panes = "A2"

    # ------------------------------------------------------------------
    # 6. Set column widths
    # ------------------------------------------------------------------
    for col_letter, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    # ------------------------------------------------------------------
    # 7. Serialise to BytesIO and return
    # ------------------------------------------------------------------
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
